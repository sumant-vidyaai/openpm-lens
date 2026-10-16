#!/usr/bin/env python3
"""Builds the openpm-lens fill-in template (Excel + CSV) and a made-up example.

Run:  python3 scripts/make_template.py
Output: packages/web/public/templates/
Column names MUST match packages/core/src/template.ts (TEMPLATE_COLUMNS) -- a test checks this.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(__file__), '..', 'packages', 'web', 'public', 'templates')
COLUMNS = ['Project', 'Programme', 'Team', 'ID', 'Task', 'Type', 'Status', 'Assignee',
           'Planned start', 'Planned finish', 'Actual start', 'Actual finish',
           'Estimate (h)', 'Actual (h)', '% complete', 'Created', 'Severity', 'Found in']
DATE_COLS = ['Planned start', 'Planned finish', 'Actual start', 'Actual finish', 'Created']
STATUSES = ['Not started', 'In progress', 'Done', 'Cancelled']
TYPES = ['Task', 'Bug', 'Story', 'Milestone', 'Risk', 'Change']
SEVERITIES = ['Critical', 'High', 'Medium', 'Low']
RULES = ['Use the % complete column', '0/100', '50/50', 'Status-based']
DAYS = ['Weekdays', 'All days']

HELP = {
    'Project': ('Optional', 'Name of the project. Fill it in to see several projects side by side on the Portfolio tab.'),
    'Programme': ('Optional', 'The programme the project belongs to (used to group projects).'),
    'Team': ('Optional', 'The team that owns the project.'),
    'ID': ('Recommended', 'A short unique reference such as T-1. Used to point at a task when something needs fixing.'),
    'Task': ('Required', 'What the task is called.'),
    'Type': ('Optional', 'Pick from the list. Needed for the Quality tab: tasks of type Bug are counted as defects.'),
    'Status': ('Required', 'Pick from the list: Not started, In progress, Done or Cancelled.'),
    'Assignee': ('Optional', 'Who does the work. Needed for Resources and Cost in money. Separate several people with a semicolon (;).'),
    'Planned start': ('For schedule', 'Year-month-day, for example 2026-03-02. The plan, not what happened.'),
    'Planned finish': ('For schedule', 'Year-month-day. The plan, not what happened.'),
    'Actual start': ('For flow', 'The day work really began. Leave blank if it has not begun.'),
    'Actual finish': ('For flow', 'The day work really finished. Leave blank if it is not finished.'),
    'Estimate (h)': ('For cost and schedule', 'Planned effort in HOURS. Plain number. Leave blank if there is no estimate (blank is not the same as 0).'),
    'Actual (h)': ('For cost', 'Hours really spent so far. Plain number. 0 means "none spent"; blank means "not recorded".'),
    '% complete': ('For cost and schedule', 'A number from 0 to 100 (write 40, not 40%). Done tasks should be 100.'),
    'Created': ('For quality', 'The day the item was raised. Year-month-day.'),
    'Severity': ('Optional', 'For bugs: Critical, High, Medium or Low.'),
    'Found in': ('Optional', 'Where a bug was found, such as Testing or Production. Lets the Quality tab count problems found after release.'),
}

# A made-up example. Every number is invented.
EXAMPLE = [
    # Project, Programme, Team, ID, Task, Type, Status, Assignee, PS, PF, AS, AF, Est, Act, Pct, Created, Sev, Found
    ['Website relaunch', 'Digital', 'Web', 'W-1', 'Requirements', 'Task', 'Done', 'Asha', '2026-03-02', '2026-03-06', '2026-03-02', '2026-03-06', 40, 44, 100, '', '', ''],
    ['Website relaunch', 'Digital', 'Web', 'W-2', 'Design mock-ups', 'Task', 'Done', 'Ravi', '2026-03-09', '2026-03-13', '2026-03-09', '2026-03-16', 40, 52, 100, '', '', ''],
    ['Website relaunch', 'Digital', 'Web', 'W-3', 'Build pages', 'Story', 'In progress', 'Asha', '2026-03-16', '2026-04-03', '2026-03-17', '', 80, 56, 60, '', '', ''],
    ['Website relaunch', 'Digital', 'Web', 'W-4', 'Build search', 'Story', 'In progress', 'Meena', '2026-03-16', '2026-03-27', '2026-03-18', '', 48, 30, 50, '', '', ''],
    ['Website relaunch', 'Digital', 'Web', 'W-5', 'Content migration', 'Task', 'Not started', 'Ravi', '2026-04-06', '2026-04-17', '', '', 60, 0, 0, '', '', ''],
    ['Website relaunch', 'Digital', 'Web', 'W-6', 'Broken contact form', 'Bug', 'Done', 'Meena', '', '', '2026-03-10', '2026-03-12', 8, 6, 100, '2026-03-09', 'High', 'Testing'],
    ['Website relaunch', 'Digital', 'Web', 'W-7', 'Logo blurry on phones', 'Bug', 'In progress', 'Asha', '', '', '2026-03-24', '', 4, 2, 50, '2026-03-23', 'Low', 'Production'],
    ['Website relaunch', 'Digital', 'Web', 'W-8', 'Launch', 'Milestone', 'Not started', 'Asha', '2026-04-20', '2026-04-20', '', '', 8, 0, 0, '', '', ''],
    ['Mobile app', 'Digital', 'Apps', 'M-1', 'Login screen', 'Story', 'Done', 'Dev', '2026-03-02', '2026-03-13', '2026-03-02', '2026-03-12', 60, 54, 100, '', '', ''],
    ['Mobile app', 'Digital', 'Apps', 'M-2', 'Profile screen', 'Story', 'Done', 'Dev', '2026-03-09', '2026-03-20', '2026-03-13', '2026-03-24', 50, 62, 100, '', '', ''],
    ['Mobile app', 'Digital', 'Apps', 'M-3', 'Payments', 'Story', 'In progress', 'Kiran', '2026-03-16', '2026-04-10', '2026-03-16', '', 120, 70, 40, '', '', ''],
    ['Mobile app', 'Digital', 'Apps', 'M-4', 'Crash on start', 'Bug', 'Done', 'Kiran', '', '', '2026-03-20', '2026-03-21', 6, 9, 100, '2026-03-19', 'Critical', 'Production'],
    ['Mobile app', 'Digital', 'Apps', 'M-5', 'Notifications', 'Story', 'Not started', 'Dev', '2026-04-06', '2026-04-24', '', '', 70, 0, 0, '', '', ''],
    ['Mobile app', 'Digital', 'Apps', 'M-6', 'Store listing', 'Task', 'Cancelled', 'Kiran', '2026-03-30', '2026-04-03', '', '', 16, 0, 0, '', '', ''],
]
EXAMPLE_SETTINGS = {'Status date': '2026-03-31', 'Currency': 'INR', 'Default hourly rate': 40}
EXAMPLE_RATES = [('Asha', 55), ('Ravi', 50), ('Meena', 45), ('Dev', 60), ('Kiran', 65)]

HEAD_FILL = PatternFill('solid', fgColor='0B6B73')
HEAD_FONT = Font(bold=True, color='FFFFFF')
INPUT_FILL = PatternFill('solid', fgColor='FFF8DC')
THIN = Side(style='thin', color='BBBBBB')


def build(path, example):
    wb = Workbook()
    # --- Start here
    s = wb.active
    s.title = 'Start here'
    lines = [
        ('openpm-lens template', 'title'),
        ('Fill in the Tasks sheet, optionally the Settings and Rates sheets, then upload this file to openpm-lens (Import or Portfolio tab). The app recognises it and sets itself up.', None),
        ('Your file stays on your computer: the app runs in your browser and sends nothing anywhere.', None),
        ('', None),
        ('Rules that matter', 'h'),
        ('1. One row per task. Do not rename, move or delete the column headings in row 1 of the Tasks sheet.', None),
        ('2. Effort is in HOURS, as plain numbers. Percent is 0 to 100.', None),
        ('3. Dates are year-month-day, for example 2026-03-31, typed as text so Excel will not change them.', None),
        ('4. A blank cell means "not known". It is never treated as 0. Type 0 only when you mean zero.', None),
        ('5. Fill only what you have. Each part of the app tells you what it still needs.', None),
        ('6. Yellow cells are for you to fill in. Cells with a small arrow have a list to pick from.', None),
        ('', None),
        ('The Tasks columns', 'h'),
    ]
    r = 1
    for text, kind in lines:
        c = s.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'title':
            c.font = Font(bold=True, size=16)
        elif kind == 'h':
            c.font = Font(bold=True, size=12)
        r += 1
    for j, h in enumerate(['Column', 'Needed for', 'What to put there'], start=1):
        c = s.cell(row=r, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    r += 1
    for col in COLUMNS:
        need, what = HELP[col]
        s.cell(row=r, column=1, value=col).font = Font(bold=True)
        s.cell(row=r, column=2, value=need)
        w = s.cell(row=r, column=3, value=what)
        w.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    s.column_dimensions['A'].width = 22
    s.column_dimensions['B'].width = 24
    s.column_dimensions['C'].width = 90
    for row in range(1, 14):
        s.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        if row in (2, 3, 6, 7, 8, 9, 10, 11):
            s.row_dimensions[row].height = 32 if row in (2, 6, 9, 10) else 18

    # --- Tasks
    t = wb.create_sheet('Tasks')
    for j, h in enumerate(COLUMNS, start=1):
        c = t.cell(row=1, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical='center')
        t.column_dimensions[get_column_letter(j)].width = max(12, len(h) + 4)
    t.column_dimensions['E'].width = 30
    t.row_dimensions[1].height = 30
    t.freeze_panes = 'F2'
    MAXR = 2000
    STYLED = 500  # kept small: a sheet over ~512 KB unpacked would need a web worker, which the app's security policy blocks
    rows = EXAMPLE if example else []
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            if v != '':
                t.cell(row=i, column=j, value=v)
    for i in range(2, STYLED + 1):
        for j, h in enumerate(COLUMNS, start=1):
            c = t.cell(row=i, column=j)
            c.fill = INPUT_FILL
            if h in DATE_COLS or h in ('Project', 'Programme', 'Team', 'ID', 'Task', 'Assignee', 'Found in'):
                c.number_format = '@'
    L = lambda h: get_column_letter(COLUMNS.index(h) + 1)

    def listval(h, values, title):
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle=title, error='Pick one from the list.')
        t.add_data_validation(dv)
        dv.add(f'{L(h)}2:{L(h)}{MAXR}')

    listval('Type', TYPES, 'Type')
    listval('Status', STATUSES, 'Status')
    listval('Severity', SEVERITIES, 'Severity')
    for h in DATE_COLS:
        col = L(h)
        dv = DataValidation(type='custom',
                            formula1=f'OR({col}2="",AND(LEN({col}2)=10,MID({col}2,5,1)="-",MID({col}2,8,1)="-",ISNUMBER(--LEFT({col}2,4)),ISNUMBER(--MID({col}2,6,2)),ISNUMBER(--RIGHT({col}2,2))))',
                            allow_blank=True, showErrorMessage=True, errorTitle='Date',
                            error='Write the date as year-month-day, for example 2026-03-31.')
        t.add_data_validation(dv)
        dv.add(f'{col}2:{col}{MAXR}')
    for h in ('Estimate (h)', 'Actual (h)'):
        dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                            showErrorMessage=True, errorTitle='Hours', error='Enter hours as a plain number, 0 or more.')
        t.add_data_validation(dv)
        dv.add(f'{L(h)}2:{L(h)}{MAXR}')
    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='100', allow_blank=True,
                        showErrorMessage=True, errorTitle='% complete', error='Enter a number from 0 to 100 (write 40, not 40%).')
    t.add_data_validation(dv)
    dv.add(f'{L("% complete")}2:{L("% complete")}{MAXR}')

    # --- Settings
    g = wb.create_sheet('Settings')
    for j, h in enumerate(['Setting', 'Value', 'What it does'], start=1):
        c = g.cell(row=1, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    settings = [
        ('Status date', '', 'The "as of" day for the report, year-month-day. Required for schedule, flow and quality. Never assumed.'),
        ('Hours per working day', 8, 'Used to turn days into hours and to spread planned work over days.'),
        ('Days counted', 'Weekdays', 'Weekdays (Mon-Fri; holidays are not known) or All days.'),
        ('Weekly capacity (h)', 40, 'Hours one person can work in a week, for the Resources tab.'),
        ('Currency', '', 'A label such as INR or USD. Optional; for the Cost in money tab.'),
        ('Default hourly rate', '', 'Rate used for people not listed on the Rates sheet. Optional.'),
        ('Green from', 0.95, 'CPI/SPI at or above this is Green. A starting value, not a standard.'),
        ('Amber from', 0.85, 'CPI/SPI at or above this (but below Green) is Amber; below is Red. A starting value, not a standard.'),
        ('Progress rule', RULES[0], 'How % complete is decided. The first choice uses the % complete column.'),
        ('In-progress credit (%)', 50, 'Only for the Status-based rule: % credited to tasks that are In progress.'),
    ]
    for i, (k, v, d) in enumerate(settings, start=2):
        if example and k in EXAMPLE_SETTINGS:
            v = EXAMPLE_SETTINGS[k]
        g.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = g.cell(row=i, column=2, value=v if v != '' else None)
        c.fill = INPUT_FILL
        if k == 'Status date':
            c.number_format = '@'
        g.cell(row=i, column=3, value=d).alignment = Alignment(wrap_text=True, vertical='top')
    g.column_dimensions['A'].width = 26
    g.column_dimensions['B'].width = 28
    g.column_dimensions['C'].width = 95
    dv = DataValidation(type='list', formula1='"' + ','.join(DAYS) + '"', allow_blank=True)
    g.add_data_validation(dv); dv.add('B4')
    dv = DataValidation(type='list', formula1='"' + ','.join(RULES) + '"', allow_blank=True)
    g.add_data_validation(dv); dv.add('B10')

    # --- Rates
    rt = wb.create_sheet('Rates')
    for j, h in enumerate(['Person', 'Hourly rate'], start=1):
        c = rt.cell(row=1, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for i in range(2, 202):
        for j in (1, 2):
            rt.cell(row=i, column=j).fill = INPUT_FILL
    if example:
        for i, (p, v) in enumerate(EXAMPLE_RATES, start=2):
            rt.cell(row=i, column=1, value=p)
            rt.cell(row=i, column=2, value=v)
    rt.column_dimensions['A'].width = 28
    rt.column_dimensions['B'].width = 16
    rt.cell(row=1, column=4, value='Names must match the Assignee column exactly. Optional.')
    wb.active = 1
    wb.save(path)


def write_csv(path, example):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        if example:
            w.writerows(EXAMPLE)


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    build(os.path.join(OUT, 'openpm-lens-template.xlsx'), False)
    build(os.path.join(OUT, 'openpm-lens-example.xlsx'), True)
    write_csv(os.path.join(OUT, 'openpm-lens-template.csv'), False)
    write_csv(os.path.join(OUT, 'openpm-lens-example.csv'), True)
    print('written to', os.path.normpath(OUT))
